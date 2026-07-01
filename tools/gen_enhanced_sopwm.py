#!/usr/bin/env python3
"""Generate enhanced_sopwm.c / .h from enhanced.xlsx."""

import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "enhanced.xlsx"
OUT_C = ROOT / "pwm" / "sopwm" / "enhanced_sopwm.c"
OUT_H = ROOT / "pwm" / "sopwm" / "enhanced_sopwm.h"

SHEETS = [
    ("M=3", 7, 3, "ENHANCED_SOPWM_LUT_N7"),
    ("M=4", 9, 4, "ENHANCED_SOPWM_LUT_N9"),
    ("M=5", 11, 5, "ENHANCED_SOPWM_LUT_N11"),
    ("M=6", 13, 6, "ENHANCED_SOPWM_LUT_N13"),
    ("M=7", 15, 7, "ENHANCED_SOPWM_LUT_N15"),
    ("M=10", 21, 10, "ENHANCED_SOPWM_LUT_N21"),
]


def gen_lut_block(ws, n_angles, lut_name):
    lines = [f"const float {lut_name}[ENHANCED_SOPWM_M_COUNT][{n_angles}] = {{"]
    for r in range(1, 101):
        m = ws.cell(r, 2).value
        angles = [ws.cell(r, c).value for c in range(3, 3 + n_angles)]
        if any(a is None for a in angles):
            raise ValueError(f"row {r} missing angle data")
        ang_str = ", ".join(f"{float(a):.10f}f" for a in angles)
        lines.append(f"    /* m={float(m):.2f} */ {{{ang_str}}},")
    lines[-1] = lines[-1].rstrip(",")
    lines.append("};")
    return "\n".join(lines)


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    h_lines = [
        "/*",
        " *   enhanced_sopwm.h - Enhanced SOPWM Look-Up Table Declarations",
        " *   Source data: enhanced.xlsx",
        " */",
        "",
        "#ifndef _ENHANCED_SOPWM_H_",
        "#define _ENHANCED_SOPWM_H_",
        "",
        "#include <stdint.h>",
        "",
        "#define ENHANCED_SOPWM_M_COUNT   100",
        "",
    ]
    for _, _n, m, name in SHEETS:
        h_lines.append(f"extern const float {name}[ENHANCED_SOPWM_M_COUNT][{m}];")
    h_lines += ["", "#endif /* _ENHANCED_SOPWM_H_ */", ""]

    c_lines = [
        "/*",
        " *   enhanced_sopwm.c - Enhanced SOPWM Look-Up Table Definitions",
        " *   Source data: enhanced.xlsx (sheets M=3..M=7, M=10)",
        " */",
        "",
        '#include "enhanced_sopwm.h"',
        "",
    ]
    for sheet, n, m, name in SHEETS:
        c_lines.append(f"/* N={n}, M={m} — sheet \"{sheet}\" */")
        c_lines.append(gen_lut_block(wb[sheet], m, name))
        c_lines.append("")

    OUT_H.write_text("\n".join(h_lines), encoding="utf-8")
    OUT_C.write_text("\n".join(c_lines), encoding="utf-8")
    print(f"Wrote {OUT_C} ({OUT_C.stat().st_size} bytes)")
    print(f"Wrote {OUT_H} ({OUT_H.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
